import sys, os, re, json, hashlib, shutil, subprocess, traceback
from pathlib import Path

import fitz
from pptx import Presentation

try:
    import pytesseract
    from PIL import Image
except Exception:
    pytesseract = None
    Image = None

source = Path(sys.argv[1]).resolve()
out = Path(sys.argv[2]).resolve()

if out.exists():
    shutil.rmtree(out)
(out / "documents").mkdir(parents=True, exist_ok=True)
(out / "pages").mkdir(parents=True, exist_ok=True)

inventory = []
hashes = {}
errors = []

def norm(s):
    if not s:
        return ""
    return (
        s.replace("ي","ی").replace("ى","ی").replace("ك","ک")
         .replace("\u200f","").replace("\ufeff","")
         .replace("\r\n","\n").replace("\r","\n")
    )

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            block = f.read(1024 * 1024)
            if not block:
                break
            h.update(block)
    return h.hexdigest()

def safe_id(path, digest):
    stem = re.sub(r'[^A-Za-z0-9._-]+', '_', path.name)[:80]
    return f"{digest[:16]}__{stem}"

def find_tesseract():
    if pytesseract is None:
        return False
    candidates = [
        shutil.which("tesseract"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for c in candidates:
        if c and Path(c).exists():
            pytesseract.pytesseract.tesseract_cmd = str(c)
            return True
    return False

HAS_TESSERACT = find_tesseract()

def ocr_page(page):
    if not HAS_TESSERACT or Image is None:
        return None
    pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    try:
        return pytesseract.image_to_string(img, lang="fas+eng")
    except Exception:
        try:
            return pytesseract.image_to_string(img, lang="eng")
        except Exception:
            return None

def extract_pdf(path):
    pages = []
    pdf = fitz.open(path)
    for i, page in enumerate(pdf):
        txt = norm(page.get_text("text"))
        used_ocr = False
        if len(re.sub(r'\s+', '', txt)) < 35:
            ot = ocr_page(page)
            if ot:
                ot = norm(ot)
                if len(re.sub(r'\s+', '', ot)) > len(re.sub(r'\s+', '', txt)):
                    txt = ot
                    used_ocr = True
        pages.append({"page": i + 1, "text": txt, "ocr": used_ocr})
    pdf.close()
    return pages

def extract_pptx(path):
    prs = Presentation(path)
    pages = []
    for idx, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
        try:
            notes = slide.notes_slide
            for shape in notes.shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    parts.append("[NOTES]\n" + shape.text)
        except Exception:
            pass
        pages.append({"page": idx, "text": norm("\n".join(parts)), "ocr": False})
    return pages

def convert_old_ppt(path, docid):
    temp = out / "_ppt-convert" / docid
    temp.mkdir(parents=True, exist_ok=True)

    # Prefer LibreOffice if available.
    lo = shutil.which("soffice") or shutil.which("libreoffice")
    if lo:
        subprocess.run(
            [lo, "--headless", "--convert-to", "pptx", "--outdir", str(temp), str(path)],
            check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE
        )
        result = list(temp.glob("*.pptx"))
        if result:
            return result[0]

    # Fallback to installed Microsoft PowerPoint via PowerShell COM.
    ps = shutil.which("powershell") or shutil.which("pwsh")
    if ps:
        target = temp / (path.stem + ".pptx")
        escaped_src = str(path).replace("'", "''")
        escaped_dst = str(target).replace("'", "''")
        cmd = (
            "$ppt=New-Object -ComObject PowerPoint.Application;"
            "$ppt.Visible=-1;"
            f"$p=$ppt.Presentations.Open('{escaped_src}',0,0,0);"
            f"$p.SaveAs('{escaped_dst}',24);"
            "$p.Close();$ppt.Quit()"
        )
        subprocess.run([ps, "-NoProfile", "-Command", cmd], check=True)
        if target.exists():
            return target

    raise RuntimeError("Old .ppt needs LibreOffice or Microsoft PowerPoint for conversion.")

files = sorted(
    p for p in source.rglob("*")
    if p.is_file() and p.suffix.lower() in {".pdf", ".pptx", ".ppt"}
)

for idx, path in enumerate(files, 1):
    digest = sha256(path)
    rel = path.relative_to(source.parent).as_posix()
    hashes.setdefault(digest, []).append(rel)
    docid = safe_id(path, digest)

    rec = {
        "path": rel,
        "filename": path.name,
        "extension": path.suffix.lower(),
        "size": path.stat().st_size,
        "sha256": digest,
        "doc_id": docid,
    }

    try:
        # Only extract canonical bytes once; duplicate paths remain in the manifest.
        if len(hashes[digest]) > 1:
            rec["status"] = "DUPLICATE_SKIPPED"
            inventory.append(rec)
            print(f"[{idx}/{len(files)}] DUPLICATE {rel}")
            continue

        if path.suffix.lower() == ".pdf":
            pages = extract_pdf(path)
        elif path.suffix.lower() == ".pptx":
            pages = extract_pptx(path)
        else:
            converted = convert_old_ppt(path, docid)
            pages = extract_pptx(converted)

        rec["page_count"] = len(pages)
        rec["ocr_pages"] = sum(1 for p in pages if p["ocr"])
        rec["text_chars"] = sum(len(p["text"]) for p in pages)
        rec["status"] = "OK"

        with open(out / "pages" / f"{docid}.jsonl", "w", encoding="utf-8") as f:
            for page in pages:
                f.write(json.dumps(page, ensure_ascii=False) + "\n")

        full_text = "\n\n".join(
            f"===== PAGE/SLIDE {p['page']} =====\n{p['text']}" for p in pages
        )
        (out / "documents" / f"{docid}.txt").write_text(full_text, encoding="utf-8")

    except Exception as e:
        rec["status"] = "FAILED"
        rec["error"] = repr(e)
        errors.append({
            "path": rel,
            "stage": "extract",
            "error": repr(e),
            "trace": traceback.format_exc(),
        })

    inventory.append(rec)
    print(f"[{idx}/{len(files)}] {rec['status']} {rel}")

duplicates = []
for digest, paths in hashes.items():
    if len(paths) > 1:
        duplicates.append({
            "sha256": digest,
            "canonical": paths[0],
            "duplicates": paths[1:],
            "all_paths": paths,
        })

with open(out / "inventory.jsonl", "w", encoding="utf-8") as f:
    for item in inventory:
        f.write(json.dumps(item, ensure_ascii=False) + "\n")

(out / "duplicate-manifest.json").write_text(
    json.dumps(duplicates, ensure_ascii=False, indent=2), encoding="utf-8"
)
(out / "errors.json").write_text(
    json.dumps(errors, ensure_ascii=False, indent=2), encoding="utf-8"
)

summary = {
    "total_entries": len(inventory),
    "canonical_documents": sum(x.get("status") == "OK" for x in inventory),
    "duplicates_skipped": sum(x.get("status") == "DUPLICATE_SKIPPED" for x in inventory),
    "failed": sum(x.get("status") == "FAILED" for x in inventory),
    "unique_sha256": len(hashes),
    "duplicate_groups": len(duplicates),
    "total_text_chars": sum(x.get("text_chars", 0) for x in inventory),
    "ocr_pages": sum(x.get("ocr_pages", 0) for x in inventory),
    "tesseract_available": HAS_TESSERACT,
}

(out / "extraction-summary.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
)

manifest = {
    "format": "sakhtyar-knowledge-corpus-export",
    "format_version": "1.0",
    "source_path": str(source),
    "summary": summary,
    "contains_ai_output": False,
    "contains_embeddings": False,
}
(out / "manifest.json").write_text(
    json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
)

print(json.dumps(summary, ensure_ascii=False, indent=2))