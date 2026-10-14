from __future__ import annotations
from pathlib import Path as _SakhtYarPath
import sys as _sakhtyar_sys
_SAKHTYAR_PYTHON_SRC = _SakhtYarPath(__file__).resolve().parents[2]
if str(_SAKHTYAR_PYTHON_SRC) not in _sakhtyar_sys.path:
    _sakhtyar_sys.path.insert(0, str(_SAKHTYAR_PYTHON_SRC))
import argparse, base64, os, re, shutil, subprocess, tempfile, urllib.request
from pathlib import Path

FAS_URL = "https://raw.githubusercontent.com/tesseract-ocr/tessdata_fast/main/fas.traineddata"

def b64(s: str) -> str:
    return base64.b64encode(s.encode("utf-8")).decode("ascii")

def clean(s: str) -> str:
    s = s.replace("\x00", " ").replace("\u200f", " ").replace("\u200e", " ")
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", s)).strip()

def emit(page: int, method: str, text: str):
    text = clean(text)
    quality = min(1.0, len(text.strip()) / 400.0)
    print(f"{page}\t{method}\t{quality:.5f}\t{b64(text)}")

def native(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            for index, page in enumerate(doc, 1):
                emit(index, "PYMUPDF_NATIVE", page.get_text("text") or "")
    elif ext == ".docx":
        from docx import Document
        doc = Document(path)
        emit(1, "PYTHON_DOCX", "\n".join(p.text for p in doc.paragraphs))
    elif ext == ".pptx":
        from pptx import Presentation
        deck = Presentation(path)
        for index, slide in enumerate(deck.slides, 1):
            emit(index, "PYTHON_PPTX", "\n".join(s.text for s in slide.shapes if hasattr(s, "text") and s.text))
    elif ext in (".txt", ".md"):
        emit(1, "TEXT_NATIVE", path.read_text(encoding="utf-8", errors="replace"))
    else:
        raise RuntimeError("unsupported file: " + ext)

def ensure_fas(tessdata_path: Path, executable: str) -> Path:
    tessdata_path.mkdir(parents=True, exist_ok=True)
    target = tessdata_path / "fas.traineddata"
    if target.is_file() and target.stat().st_size > 100000:
        return target
    for source in (Path(executable).parent/"tessdata"/"fas.traineddata",
                   Path(r"C:\Program Files\Tesseract-OCR\tessdata\fas.traineddata"),
                   Path(r"C:\Program Files (x86)\Tesseract-OCR\tessdata\fas.traineddata")):
        if source.is_file():
            shutil.copy2(source, target)
            if target.stat().st_size > 100000:
                return target
    part = target.with_suffix(".traineddata.part")
    try:
        print("[OCR-BOOTSTRAP] downloading fas.traineddata to project repository", flush=True)
        urllib.request.urlretrieve(FAS_URL, part)
        if part.stat().st_size <= 100000:
            raise RuntimeError("downloaded language file is unexpectedly small")
        part.replace(target)
        return target
    except Exception as exc:
        try: part.unlink(missing_ok=True)
        except OSError: pass
        raise RuntimeError(f"Persian OCR data recovery failed: {target}; {exc}") from exc

def configure_tesseract():
    import pytesseract
    candidates=[os.environ.get("SAKHTYAR_TESSERACT_EXE"),
                r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"]
    executable=next((x for x in candidates if x and Path(x).is_file()),None)
    if not executable:
        raise RuntimeError("Tesseract OCR executable not found")
    pytesseract.pytesseract.tesseract_cmd=executable
    root=Path(__file__).resolve().parents[2]
    tessdata_path=Path(os.environ.get("SAKHTYAR_TESSDATA_DIR") or root/".sakhtyar"/"repository"/"tesseract"/"tessdata").resolve()
    ensure_fas(tessdata_path,executable)
    os.environ.pop("TESSDATA_PREFIX",None)
    return executable,tessdata_path

def is_visually_blank(image_name: str) -> bool:
    """Conservative heuristic. Blank pages are still queued for human confirmation."""
    from PIL import Image
    with Image.open(image_name) as image:
        gray = image.convert("L")
        gray.thumbnail((450, 450))
        w, h = gray.size
        inset_x, inset_y = max(1, w // 25), max(1, h // 25)
        if w <= inset_x * 2 or h <= inset_y * 2:
            return False
        inner = gray.crop((inset_x, inset_y, w - inset_x, h - inset_y))
        histogram = inner.histogram()
        return sum(histogram[:210]) <= max(2, int(inner.width * inner.height * 0.00015))


def run_tesseract_png(executable: str, image_name: str, tessdata_path: Path, lang: str) -> str:
    env = os.environ.copy()
    env.pop("TESSDATA_PREFIX", None)
    diagnostic = []
    succeeded = 0

    def attempt(filename: str, psm: str, variant: str) -> str:
        nonlocal succeeded
        command = [executable, filename, "stdout", "--tessdata-dir", str(tessdata_path),
                   "-l", lang, "--psm", psm]
        result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                env=env, check=False, timeout=90)
        if result.returncode != 0:
            detail = result.stderr.decode("utf-8", errors="replace")[:220]
            diagnostic.append(f"{variant}/{psm}: error {result.returncode}: {detail}")
            return ""
        succeeded += 1
        value = clean(result.stdout.decode("utf-8", errors="replace"))
        diagnostic.append(f"{variant}/{psm}: chars={len(value)}")
        return value

    choices = [attempt(image_name, psm, "native-raster") for psm in ("3", "6", "11")]
    choices = [x for x in choices if x]
    if choices:
        return max(choices, key=len)

    from PIL import Image, ImageOps, ImageEnhance
    with Image.open(image_name) as image:
        enhanced = ImageEnhance.Contrast(ImageOps.autocontrast(image.convert("L"), cutoff=1)).enhance(1.6)
        binary = enhanced.point(lambda pixel: 255 if pixel >= 195 else 0)
        for name, variant in (("enhanced", enhanced), ("binary", binary)):
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as t:
                tmp = t.name
            try:
                variant.save(tmp)
                for psm in ("6", "11"):
                    value = attempt(tmp, psm, name)
                    if value:
                        return value
            finally:
                try: os.unlink(tmp)
                except OSError: pass

    # Tesseract failure (bad binary/language pack) must not be disguised as an unreadable page.
    if succeeded == 0:
        raise RuntimeError("Tesseract executable/configuration failure: " + " | ".join(diagnostic))
    return ""  # Successful OCR invocations with zero text => human review, not pipeline failure.


def ocr_pdf(path: Path, pages: list[int], lang: str):
    import pymupdf
    executable, tessdata_path = configure_tesseract()
    with pymupdf.open(path) as doc:
        for page_no in pages:
            if not 1 <= page_no <= len(doc):
                raise RuntimeError(f"OCR page out of range: {path.name} page={page_no}")
            page = doc[page_no - 1]
            pix = page.get_pixmap(matrix=pymupdf.Matrix(3, 3), alpha=False)
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as t:
                image_name = t.name
            try:
                pix.save(image_name)
                if is_visually_blank(image_name) and not clean(page.get_text("text") or ""):
                    emit(page_no, "OCR_REVIEW_BLANK", "")
                    continue
                text = run_tesseract_png(executable, image_name, tessdata_path, lang)
                if not text:
                    emit(page_no, "OCR_REVIEW_UNREADABLE", "")
                else:
                    emit(page_no, "TESSERACT_" + lang, text)
            finally:
                try: os.unlink(image_name)
                except OSError: pass

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--mode",choices=["native","ocr"],required=True)
    parser.add_argument("--file",required=True)
    parser.add_argument("--pages",default="")
    parser.add_argument("--lang",default="fas")
    args=parser.parse_args()
    path=Path(args.file)
    if args.mode=="native": native(path)
    else: ocr_pdf(path,[int(x) for x in args.pages.split(",") if x.strip()],args.lang)

if __name__=="__main__":
    main()
